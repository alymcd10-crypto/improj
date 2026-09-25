"""SYNTHETIC test data. Every name, email and phone number here is fictional
(555-01xx numbers and example.com/.net addresses are reserved for fiction).
No client data is used anywhere in the tests.
"""
from __future__ import annotations

import csv
import io

TP_HEADERS = ["Contact ID", "First Name", "Last Name", "Spouse First Name", "Spouse Last Name", "Company",
              "Email", "Email 2", "Home Phone", "Mobile Phone", "Work Phone", "Work Phone Ext", "Address",
              "City", "State", "Zip", "Contact Type", "Notes", "Last Modified"]


def tp(cid, first="", last="", sf="", sl="", company="", email="", email2="", home="", mobile="", work="",
       work_ext="", address="", city="Seattle", state="WA", zip_="98101", types="", notes="",
       modified="2026-09-01"):
    return dict(zip(TP_HEADERS, [cid, first, last, sf, sl, company, email, email2, home, mobile, work, work_ext,
                                 address, city, state, zip_, types, notes, modified]))


def base_contacts() -> list[dict]:
    return [
        tp("T001", "Avery", "Sample", email="avery@example.com", mobile="206-555-0101",
           address="12 Test Street", zip_="02134", types="Past Client; Sphere"),
        tp("T002", "Avery", "Sample", email="avery.s@example.net", mobile="425-555-0102",
           address="900 Other Ave"),
        tp("T003", "Jordan", "Example", "Casey", "Example", email="example.family@example.com",
           home="206-555-0103", address="5 Shared Lane", types="Sphere"),
        tp("T004", "Morgan", "Example", email="example.family@example.com", home="206-555-0103",
           address="5 Shared Lane"),
        tp("T005", company="Acme Test Holdings LLC", email="office@acmetest.example", work="206-555-0105",
           work_ext="12", types="Vendor"),
        tp("T006", "Renée", "O'Testly", email="renee@example.com", mobile="(206) 555-0106",
           notes="Kids: Emma and Liam.\nPrefers text messages.", types="Buyer"),
        tp("T007", "Sam", "Newerton", mobile="206-555-0107",
           notes="2026-08-01 client said old cell 206-555-0177 is disconnected; new cell on file."),
        tp("T008", "Pat", "Optout", email="pat@example.com", types="Past Client",
           notes="2026-07-15 asked to opt out of the newsletter."),
        tp("T009", "Lee", "Verified", email="lee@example.com", mobile="206-555-0109"),
        tp("T010", "Chris", "Categories", email="chris@example.com", types="Sphere; Past Client"),
        tp("T011", "Kim", "Unmapped", email="kim@example.com", types="Past Client"),
        tp("T012", "Quinn", "Formula", email="quinn@example.com", notes="=HYPERLINK(\"http://bad.example\")"),
        tp("T013", "Taylor", "Extension", email="taylor@example.com", work="206-555-0113", work_ext="45"),
        tp("T014", "Jamie", "Missing", email="jamie@example.com", mobile="206-555-0114"),
        tp("T015", "Drew", "Notes", email="drew@example.com",
           notes="Birthday party 6/12.\n\"Loves   BOATING\" — mentioned at closing, 2019."),
        tp("T016", "Blair", "Longnote", email="blair@example.com", notes="Short note."),
    ]


WL_HEADERS = ["Change_ID", "Queue", "Status", "Household", "First_Name", "Last_Name", "Company", "Match_Email",
              "Match_Other_Email", "Match_Phones", "Match_Address", "Match_City", "Match_Zip", "TP_ID", "Field",
              "Action", "Source_Value", "Source_Ref", "Proposed_Addition", "Reason", "Execution_Rule",
              "Identity_Rule", "Current_TP_Value", "Actual_Change", "Verified_After_Value", "Verified_At_UTC",
              "Operator", "Run_ID", "Result_Note"]


def wl(cid, action, proposed, first="", last="", household="", company="", email="", phones="", address="",
       zip_="", tp_id="", queue="CORRECTION_CANDIDATE", status="PENDING_CURRENT_CHECK", field="",
       source_value="", verified_after="", result_note=""):
    row = {h: "" for h in WL_HEADERS}
    row.update({"Change_ID": cid, "Queue": queue, "Status": status, "Household": household, "First_Name": first,
                "Last_Name": last, "Company": company, "Match_Email": email, "Match_Phones": phones,
                "Match_Address": address, "Match_Zip": zip_, "TP_ID": tp_id,
                "Field": field or {"ADD_MISSING_PHONE": "Phone", "ADD_MISSING_CATEGORIES": "Contact Type",
                                   "RESTORE_MISSING_NOTE_DETAIL": "Notes"}.get(action, action),
                "Action": action, "Source_Value": source_value or proposed, "Source_Ref": "All Contacts row 99",
                "Proposed_Addition": proposed, "Reason": "SYNTHETIC test row",
                "Execution_Rule": "Add only.", "Identity_Rule": "Name plus one identifier.",
                "Verified_After_Value": verified_after, "Result_Note": result_note})
    return row


def base_worklist() -> list[dict]:
    return [
        # Same name, disambiguated by email -> T001; category already present in different order
        wl("C001", "ADD_MISSING_CATEGORIES", "Sphere;Past Client", "Avery", "Sample", email="avery@example.com"),
        # Same name, only a name -> two candidates, name alone never matches
        wl("C002", "ADD_MISSING_PHONE", "206-555-0199", "Avery", "Sample"),
        # Couple sharing a profile, matched on household + shared phone; the phone is present in other format
        wl("C003", "ADD_MISSING_PHONE", "(206) 555.0103", household="Jordan & Casey Example",
           phones="206-555-0103"),
        # Shared family email on two profiles; spouse name identifies the right one -> T003 (category missing)
        wl("C004", "ADD_MISSING_CATEGORIES", "Past Client", "Casey", "Example",
           email="example.family@example.com"),
        # Business without first name
        wl("C005", "ADD_MISSING_PHONE", "206-555-0155", company="Acme Test Holdings LLC",
           email="office@acmetest.example"),
        # Accents/apostrophes: worklist spelled without accent and with curly apostrophe
        wl("C006", "RESTORE_MISSING_NOTE_DETAIL", "Dog named Biscuit", "Renee", "O’Testly",
           email="renee@example.com"),
        # Newer conflicting info: old number reported disconnected in notes
        wl("C007", "ADD_MISSING_PHONE", "206-555-0177", "Sam", "Newerton", phones="206-555-0107"),
        # Deliberate opt-out -> hold category addition
        wl("C008", "ADD_MISSING_CATEGORIES", "Newsletter", "Pat", "Optout", email="pat@example.com"),
        # Verified ID pointing to a different name
        wl("C009", "ADD_MISSING_PHONE", "206-555-0190", "Dana", "Mismatch", tp_id="T009"),
        # Categories already present in a different order
        wl("C010", "ADD_MISSING_CATEGORIES", "Past Client | Sphere", "Chris", "Categories", email="chris@example.com"),
        # Unmapped category plus a mapped missing one
        wl("C011", "ADD_MISSING_CATEGORIES", "Golf Buddies; Sphere", "Kim", "Unmapped", email="kim@example.com"),
        # Formula-looking values
        wl("C012", "RESTORE_MISSING_NOTE_DETAIL", "=cmd|' /C calc'!A0", "Quinn", "Formula", email="quinn@example.com"),
        # Same number, different extension
        wl("C013", "ADD_MISSING_PHONE", "206-555-0113 x46", "Taylor", "Extension", email="taylor@example.com"),
        # Genuinely missing phone -> Cowork candidate
        wl("C014", "ADD_MISSING_PHONE", "+1 206 555 0144", "Jamie", "Missing", email="jamie@example.com"),
        # Note detail present with different formatting
        wl("C015", "RESTORE_MISSING_NOTE_DETAIL", "loves boating", "Drew", "Notes", email="drew@example.com"),
        # REVIEW_ONLY row that is missing -> never promoted
        wl("C016", "ADD_MISSING_PHONE", "206-555-0166", "Jamie", "Missing", email="jamie@example.com",
           queue="REVIEW_ONLY"),
        # Nobody
        wl("C017", "ADD_MISSING_PHONE", "206-555-0170", "Nobody", "Here", email="nobody@example.com"),
        # Ambiguous digits (no extension marker)
        wl("C018", "ADD_MISSING_PHONE", "206555011455", "Jamie", "Missing", email="jamie@example.com"),
        # Whole-note replacement proposal -> review, not replacement
        wl("C019", "RESTORE_MISSING_NOTE_DETAIL", " ".join(["Long historical note sentence"] * 25), "Blair",
           "Longnote", email="blair@example.com"),
        # Unsupported / out of scope change type
        wl("C020", "REPLACE_PRIMARY_PHONE", "206-555-0101", "Avery", "Sample", email="avery@example.com"),
        # Already done by Cowork -> not re-queued
        wl("C021", "ADD_MISSING_PHONE", "206-555-0121", "Jamie", "Missing", email="jamie@example.com",
           status="DONE"),
    ]


def to_csv(headers: list[str], rows: list[dict], delimiter: str = ",", encoding: str = "utf-8") -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=headers, delimiter=delimiter, lineterminator="\r\n",
                       extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode(encoding)


def export_bytes(contacts=None, drop: tuple[str, ...] = (), blank: tuple[str, ...] = (), **kw) -> bytes:
    contacts = contacts if contacts is not None else base_contacts()
    headers = [h for h in TP_HEADERS if h not in drop]
    rows = [{**c, **{b: "" for b in blank}} for c in contacts]
    return to_csv(headers, rows, **kw)


def worklist_bytes(rows=None, **kw) -> bytes:
    return to_csv(WL_HEADERS, rows if rows is not None else base_worklist(), **kw)


def mapping_rows() -> list[dict]:
    return [
        {"Source_Category": "Past Client", "TP_Contact_Type": "Past Client", "Status": "APPROVED"},
        {"Source_Category": "Sphere", "TP_Contact_Type": "Sphere", "Status": "APPROVED"},
        {"Source_Category": "Newsletter", "TP_Contact_Type": "Newsletter", "Status": "APPROVED"},
        {"Source_Category": "Golf Buddies", "TP_Contact_Type": "", "Status": "UNRESOLVED"},
    ]


def master_xlsx_bytes() -> bytes:
    """Workbook with a title row, All Contacts, an overlapping agent tab and a within-tab duplicate."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Contacts"
    ws.append(["SYNTHETIC master export"])  # title row above the header
    ws.append(["First Name", "Last Name", "Email", "Mobile Phone", "Zip", "Category", "Birthday"])
    rows = [["Avery", "Sample", "avery@example.com", "206-555-0101", 2134, "Past Client", "1800-04-02"],
            ["Renée", "O'Testly", "renee@example.com", "206-555-0106", 98101, "Buyer", "1975-01-09"],
            ["Avery", "Sample", "avery@example.com", "206-555-0101", 2134, "Sphere", ""]]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=3, min_col=5, max_col=5):
        row[0].number_format = "00000"
    ag = wb.create_sheet("Agent - Leslie")
    ag.append(["First Name", "Last Name", "Email", "Mobile Phone"])
    ag.append(["Renée", "O'Testly", "renee@example.com", "206-555-0106"])
    ag.append(["Solo", "Tabonly", "solo@example.com", "206-555-0188"])
    idn = wb.create_sheet("Identity Review")
    idn.append(["Finding", "Name", "Status"])
    idn.append(["Two profiles may be the same person", "Avery Sample", ""])
    idn.append(["Resolved earlier", "Chris Categories", "Resolved"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
